import pandas as pd
from datetime import datetime, timedelta
from tkinter import Tk, filedialog, Toplevel, Label, DoubleVar
from tkinter.ttk import Progressbar  # Importar Progressbar desde ttk
from openpyxl import load_workbook
from tqdm import tqdm  # Para la barra de progreso en consola

# ==============================
# FUNCIONES AUXILIARES
# ==============================

def seleccionar_archivo(titulo):
    """Ventana para seleccionar archivo"""
    root = Tk()
    root.withdraw()
    return filedialog.askopenfilename(title=titulo)


def leer_archivo(ruta):
    """Lee CSV o Excel"""
    if ruta.endswith(".csv"):
        return pd.read_csv(ruta, low_memory=False)
    else:
        return pd.read_excel(ruta)


def parse_fecha(fecha_str):
    """Convierte fecha sin zona horaria"""
    try:
        fecha = pd.to_datetime(fecha_str, dayfirst=True, errors='coerce')
        if fecha is not None and not pd.isna(fecha):
            if fecha.tzinfo:
                fecha = fecha.tz_localize(None)
        return fecha
    except:
        return None


def calcular_fecha_inova(pull_time, age):
    """
    Calcula fecha real en INOVA usando Age
    Ejemplo: '0d 08h 06m'
    """
    try:
        fecha_base = pd.to_datetime(pull_time)

        if fecha_base is not None and not pd.isna(fecha_base):
            if fecha_base.tzinfo:
                fecha_base = fecha_base.tz_localize(None)

        dias = horas = minutos = 0

        if isinstance(age, str):
            if "d" in age:
                dias = int(age.split("d")[0])
            if "h" in age:
                horas = int(age.split("h")[0].split()[-1])
            if "m" in age:
                minutos = int(age.split("m")[0].split()[-1])

        return fecha_base - timedelta(days=dias, hours=horas, minutes=minutos)

    except:
        return None


def mostrar_barra_progreso():
    """Crea una ventana con barra de progreso visual"""
    root = Tk()
    root.title("Procesando archivo")
    root.geometry("400x150")

    label = Label(root, text="Procesando registros...", width=20)
    label.pack(pady=10)

    barra_progreso = Progressbar(root, length=300, mode='determinate')
    barra_progreso.pack(pady=10)

    return root, barra_progreso


# ==============================
# SELECCIÓN DE ARCHIVOS
# ==============================

print("Selecciona archivo de ROBADOS")
archivo_robados = seleccionar_archivo("Selecciona archivo de ROBADOS")

print("Selecciona archivo de ESTATUS 1 (opcional)")
archivo_estatus1 = seleccionar_archivo("Selecciona estatus 1")

print("Selecciona archivo de ESTATUS 2 (opcional)")
archivo_estatus2 = seleccionar_archivo("Selecciona estatus 2")

fecha_estatus_input = input("Ingresa la fecha del estatus (dd/mm/yyyy): ")
fecha_estatus = datetime.strptime(fecha_estatus_input, "%d/%m/%Y")

# ==============================
# CARGA DE DATOS
# ==============================

df_robados = pd.read_excel(archivo_robados)

estatus_total = []

for archivo in [archivo_estatus1, archivo_estatus2]:
    if archivo:
        df = leer_archivo(archivo)
        estatus_total.append(df)

# ==============================
# NORMALIZAR ESTATUS
# ==============================

registros_estatus = []

for df in estatus_total:

    # INOVA
    if "Node" in df.columns:
        for _, row in df.iterrows():
            serie = str(row.get("Node")).strip()
            linea = row.get("Line")
            estaca = row.get("Flag")

            fecha_visto = calcular_fecha_inova(
                row.get("Pull Time"),
                row.get("Age")
            )

            registros_estatus.append({
                "serie": serie,
                "linea": linea,
                "estaca": estaca,
                "fecha": fecha_visto,
                "fuente": "INOVA"
            })

    # SERCEL
    elif "Id" in df.columns:
        for _, row in df.iterrows():
            serie = str(row.get("Id")).strip()
            linea = row.get("Topo-Location Line")
            estaca = row.get("Topo-Location Point")

            fecha_visto = pd.to_datetime(row.get("Qc Time"), errors='coerce')

            if fecha_visto is not None and not pd.isna(fecha_visto):
                if fecha_visto.tzinfo:
                    fecha_visto = fecha_visto.tz_localize(None)

            registros_estatus.append({
                "serie": serie,
                "linea": linea,
                "estaca": estaca,
                "fecha": fecha_visto,
                "fuente": "SERCEL"
            })

# ==============================
# ACTUALIZAR EXCEL (SIN PERDER FORMATO)
# ==============================

wb = load_workbook(archivo_robados)
ws = wb.active

headers = [cell.value for cell in ws[1]]

col_serie = headers.index("SERIE") + 1
col_linea = headers.index("LINEA") + 1
col_estaca = headers.index("ESTACA") + 1
col_fecha_ext = headers.index("FECHA EXTRAVIO") + 1
col_comentario = headers.index("COMENTARIO") + 1
col_estatus = headers.index("ESTATUS") + 1
col_fecha_rec = headers.index("FECHA DE RECUPERADO") + 1

# ==============================
# PROCESAMIENTO
# ==============================

# Crear la barra de progreso
root, barra_progreso = mostrar_barra_progreso()

# Definir el número total de filas a procesar
total_filas = ws.max_row - 1  # Restamos 1 por la fila de encabezados
barra_progreso['maximum'] = total_filas
barra_progreso['value'] = 0

for row in tqdm(range(2, ws.max_row + 1), total=total_filas, desc="Procesando filas", unit="fila"):
    serie_robado = str(ws.cell(row, col_serie).value).strip()
    linea_robado = ws.cell(row, col_linea).value
    estaca_robado = ws.cell(row, col_estaca).value
    fecha_ext = parse_fecha(ws.cell(row, col_fecha_ext).value)

    # Verificar si el estatus ya es "RECUPERADO", si es así, saltar a la siguiente fila
    estatus_actual = str(ws.cell(row, col_estatus).value).strip()
    if estatus_actual == "RECUPERADO":
        continue  # Saltar a la siguiente fila si ya está recuperado

    for est in registros_estatus:

        if est["serie"] == serie_robado:

            if fecha_ext is not None and est["fecha"] is not None:
                if not pd.isna(fecha_ext) and not pd.isna(est["fecha"]):

                    if fecha_ext <= est["fecha"]:

                        comentario = f"VISTO EN TENDIDO EN LR:{est['linea']} E:{est['estaca']} EL DIA {est['fecha'].strftime('%d/%m/%Y')} ({est['fuente']})"

                        ws.cell(row, col_comentario).value = comentario

                        # Si está en diferente ubicación → RECUPERADO
                        if (est["linea"] != linea_robado) or (est["estaca"] != estaca_robado):

                            ws.cell(row, col_estatus).value = "RECUPERADO"
                            ws.cell(row, col_fecha_rec).value = fecha_estatus.strftime("%d/%m/%Y")

    # Actualizar barra de progreso visual
    barra_progreso['value'] = row - 1
    barra_progreso.update()

# ==============================
# GUARDAR
# ==============================

wb.save(archivo_robados)

print("Archivo actualizado correctamente sin perder formato")
root.quit()